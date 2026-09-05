#!/usr/bin/env python3
"""build_workbook.py — write (or replace) one "<model> 对比" sheet in the master price workbook.

usage: build_workbook.py <spec.json>

spec.json
{
  "workbook": "/path/Apple_Official_Price.xlsx",      # created if missing; other sheets are preserved
  "sheet":    "MacBook Pro 对比",
  "model":    "MacBook Pro",
  "crawl_date": "2026-09-05",
  "keys": [ {"col": "chassis-dimensionScreensize", "name": "尺寸", "map": {"14inch": "14 英寸"}, "width": 10}, ... ],
  "regions": [ {"code": "US", "tsv": "/path/mbp_us_rows.tsv", "url": "https://www.apple.com/shop/buy-mac/macbook-pro"},
               {"code": "CN", "tsv": "...", "url": "..."}, {"code": "SG", "tsv": "...", "url": "..."} ],
  "reference": "CN",                                    # region every other region is compared against
  "addons": [ {"name": "Nano-texture 显示屏", "US": 150, "CN": 1125, "SG": null, "note": "..."} ],
  "notes": ["在售现货，无预购文案", "..."],             # free-text lines for the assumptions sheet
  "anomalies": ["..."]
}

Rows are matched across regions on the key columns (page radio values, identical across Apple storefronts).
All derived columns are formulas that reference the assumptions sheet, so editing a tax rate or FX rate
there recalculates every model sheet. The workbook is saved with fullCalcOnLoad so Excel computes the
formulas on open (no LibreOffice needed).
"""
import sys, csv, json, re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

ASSUME = '说明与假设'
F = 'Arial'
HDR_FILL = PatternFill('solid', fgColor='1F3864'); HDR_FONT = Font(name=F, bold=True, color='FFFFFF')
THIN = Side(style='thin', color='BFBFBF'); BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BLUE = Font(name=F, color='0000FF'); YELLOW = PatternFill('solid', fgColor='FFFF00')
PALETTE = ['E8F1FB', 'EEF7EA', 'FDF2E3', 'F3E8FB', 'FBEAEA', 'EAF7F7']

# Storefront defaults. tax_inclusive: is the listed price already tax-inclusive?  rate: local currency -> RMB.
REGION_DEFAULTS = {
    'US': dict(label='美国', currency='USD', symbol='$', tax_inclusive=False, tax=0.0625, rate=7.10,
               note='官网标价不含州销售税；默认 MA 州 6.25%（州统一税率，无地方附加，电脑属应税商品）'),
    'CN': dict(label='中国大陆', currency='CNY', symbol='¥', tax_inclusive=True, tax=0.13, rate=1.0, note='官网标价含 13% 增值税'),
    'SG': dict(label='新加坡', currency='SGD', symbol='S$', tax_inclusive=True, tax=0.09, rate=5.50, note='官网标价含 9% GST'),
    'HK': dict(label='中国香港', currency='HKD', symbol='HK$', tax_inclusive=True, tax=0.0, rate=0.91, note='无销售税'),
    'JP': dict(label='日本', currency='JPY', symbol='¥', tax_inclusive=True, tax=0.10, rate=0.048, note='官网标价含 10% 消费税'),
    'UK': dict(label='英国', currency='GBP', symbol='£', tax_inclusive=True, tax=0.20, rate=9.50, note='官网标价含 20% VAT'),
}
# fixed rows in the assumptions sheet so formulas stay stable across rebuilds
ASSUME_REGION_ROW0 = 4   # header row 3, US row 4, CN row 5, SG row 6, ...
REGION_ORDER = ['US', 'CN', 'SG', 'HK', 'JP', 'UK']

def region_row(code): return ASSUME_REGION_ROW0 + REGION_ORDER.index(code)

def ensure_assumptions(wb):
    """Create the assumptions sheet once: region table (rows 4-9) + free-text notes below."""
    if ASSUME in wb.sheetnames: return wb[ASSUME]
    ws = wb.create_sheet(ASSUME)
    ws['A1'] = 'Apple 官网价格对比 — 说明与假设'; ws['A1'].font = Font(name=F, bold=True, size=13)
    ws['A2'] = '黄色单元格为可编辑假设值（税率、汇率）；改动后所有「对比」sheet 的含税价、折算、差额、溢价自动重算。汇率为占位估算，非当日实际汇率。'
    ws['A2'].font = Font(name=F, italic=True, color='595959')
    heads = ['地区', '货币', '官网价是否含税', '税率（不含税地区用于计算含税价）', '汇率（1 本币 = ? RMB）', '口径说明']
    for c, h in enumerate(heads, 1):
        x = ws.cell(3, c, h); x.fill = HDR_FILL; x.font = HDR_FONT; x.border = BORDER; x.alignment = Alignment(horizontal='center', wrap_text=True)
    for code in REGION_ORDER:
        d = REGION_DEFAULTS[code]; r = region_row(code)
        ws.cell(r, 1, f'{code} {d["label"]}'); ws.cell(r, 2, d['currency']); ws.cell(r, 3, '是' if d['tax_inclusive'] else '否')
        t = ws.cell(r, 4, d['tax']); t.number_format = '0.00%'; t.font = BLUE; t.fill = YELLOW
        x = ws.cell(r, 5, d['rate']); x.number_format = '0.0000'; x.font = BLUE; x.fill = YELLOW
        ws.cell(r, 6, d['note'])
        for c in range(1, 7):
            ws.cell(r, c).border = BORDER
            if c not in (4, 5): ws.cell(r, c).font = Font(name=F)
    ws.cell(11, 1, '数据来源与抓取记录').font = Font(name=F, bold=True, size=12)
    for c, w in zip('ABCDEF', [16, 8, 14, 18, 18, 70]): ws.column_dimensions[c].width = w
    return ws

def append_notes(ws, model, spec):
    r = ws.max_row + 2
    ws.cell(r, 1, f'【{model}】抓取日期 {spec.get("crawl_date", "")}').font = Font(name=F, bold=True)
    for reg in spec['regions']:
        r += 1; ws.cell(r, 1, f'  {reg["code"]} 来源'); ws.cell(r, 2, reg.get('url', '')); ws.cell(r, 2).font = Font(name=F, color='0563C1')
    for line in spec.get('notes', []):
        r += 1; ws.cell(r, 1, '  说明'); ws.cell(r, 2, line); ws.cell(r, 2).alignment = Alignment(wrap_text=True)
    for line in spec.get('anomalies', []):
        r += 1; ws.cell(r, 1, '  定价异常'); ws.cell(r, 2, line); ws.cell(r, 2).alignment = Alignment(wrap_text=True)
    ws.cell(r + 1, 1, '  抓取方式'); ws.cell(r + 1, 2, '在各站定制页逐个点选全部价格维度组合并读取汇总价；总价已与「基础价 + 各项差价」核对；颜色/键盘等不改价维度验证后固定；附加项差价单列，不进乘积。')
    ws.cell(r + 1, 2).alignment = Alignment(wrap_text=True)

def colname(k, code):
    """key 'col' may be a string or {regionCode: column} when crawls named columns differently"""
    c = k['col']; return c if isinstance(c, str) else c[code]

def load_rows(tsv, keys, code):
    out = {}
    for r in csv.DictReader(open(tsv, encoding='utf-8'), delimiter='\t'):
        if str(r.get('drift', '0')) == '1': continue          # auto-switched config, recorded under its real key elsewhere
        try: amt = int(round(float(r['amount'])))
        except: continue
        out[tuple(r[colname(k, code)] for k in keys)] = amt
    return out

def sort_key(t, keys):
    out = []
    for k, v in zip(keys, t):
        order = list(k.get('map', {}).keys())
        if v in order: out.append((0, order.index(v), ''))
        else:
            m = re.match(r'^(\d+)(gb|tb)$', v)
            out.append((1, int(m.group(1)) * (1000 if m.group(2) == 'tb' else 1), '') if m else (2, 0, v))
    return out

def main(spec_path):
    spec = json.load(open(spec_path, encoding='utf-8'))
    keys = spec['keys']; regions = spec['regions']; ref = spec.get('reference', 'CN')
    codes = [r['code'] for r in regions]
    if ref not in codes: ref = codes[0]
    data = {r['code']: load_rows(r['tsv'], keys, r['code']) for r in regions}
    allkeys = sorted(set().union(*data.values()), key=lambda t: sort_key(t, keys))
    print({c: len(v) for c, v in data.items()}, 'union', len(allkeys))
    for c in codes:
        missing = [k for k in allkeys if k not in data[c]]
        if missing: print(f'  {c} lacks {len(missing)} combos, e.g. {missing[:3]}')

    try: wb = load_workbook(spec['workbook'])
    except FileNotFoundError: wb = Workbook(); wb.remove(wb.active)
    ws_a = ensure_assumptions(wb)
    if spec['sheet'] in wb.sheetnames: del wb[spec['sheet']]
    ws = wb.create_sheet(spec['sheet'], index=len(wb.sheetnames) - 1)   # keep assumptions last

    # ---- header ----
    heads = ['序号'] + [k['name'] for k in keys]
    col_of = {}  # (code, kind) -> column index
    for c in codes:
        d = REGION_DEFAULTS[c]; lab = d['label']; cur = d['currency']
        heads.append(f'{lab}官网价 ({cur})'); col_of[(c, 'list')] = len(heads)
        if not d['tax_inclusive']:
            heads.append(f'{lab}含税价 ({cur})'); col_of[(c, 'taxed')] = len(heads)
        if c != 'CN':
            heads.append(f'{lab}折算 (RMB)'); col_of[(c, 'rmb')] = len(heads)
    refd = REGION_DEFAULTS[ref]
    for c in codes:
        if c == ref: continue
        heads.append(f'差额 {refd["label"]}−{REGION_DEFAULTS[c]["label"]} (RMB)'); col_of[(c, 'diff')] = len(heads)
        heads.append(f'{refd["label"]}相对{REGION_DEFAULTS[c]["label"]}溢价 %'); col_of[(c, 'prem')] = len(heads)
    ws.append(heads)
    for c in range(1, len(heads) + 1):
        x = ws.cell(1, c); x.fill = HDR_FILL; x.font = HDR_FONT; x.border = BORDER; x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 40
    L = get_column_letter
    def rmb_expr(c, i):
        """cell expression giving region c's price in RMB for row i"""
        d = REGION_DEFAULTS[c]; rr = region_row(c)
        if c == 'CN': return f'{L(col_of[(c, "list")])}{i}'
        src = L(col_of[(c, 'taxed')]) if not d['tax_inclusive'] else L(col_of[(c, 'list')])
        return f'{src}{i}*\'{ASSUME}\'!$E${rr}'

    # ---- rows ----
    nk = len(keys); bands = {}
    for i, t in enumerate(allkeys, 2):
        bk = t[:2]; bands.setdefault(bk, PALETTE[len(bands) % len(PALETTE)])
        ws.cell(i, 1, i - 1)
        for j, (k, v) in enumerate(zip(keys, t), 2): ws.cell(i, j, k.get('map', {}).get(v, v))
        have = {c: (t in data[c]) for c in codes}
        for c in codes:
            d = REGION_DEFAULTS[c]; lc = col_of[(c, 'list')]
            if have[c]:
                ws.cell(i, lc, data[c][t])
                if not d['tax_inclusive']: ws.cell(i, col_of[(c, 'taxed')], f'={L(lc)}{i}*(1+\'{ASSUME}\'!$D${region_row(c)})')
                if c != 'CN': ws.cell(i, col_of[(c, 'rmb')], f'={rmb_expr(c, i)}')
            else:
                ws.cell(i, lc, f'{d["label"]}无此配置')
        for c in codes:
            if c == ref: continue
            if have[c] and have[ref]:
                refcell = L(col_of[(ref, 'rmb')]) + str(i) if ref != 'CN' else L(col_of[('CN', 'list')]) + str(i)
                ccell = L(col_of[(c, 'rmb')]) + str(i) if c != 'CN' else L(col_of[('CN', 'list')]) + str(i)
                ws.cell(i, col_of[(c, 'diff')], f'={refcell}-{ccell}')
                ws.cell(i, col_of[(c, 'prem')], f'=IF({ccell}=0,0,{L(col_of[(c, "diff")])}{i}/{ccell})')
        fill = PatternFill('solid', fgColor=bands[bk])
        for c in range(1, len(heads) + 1):
            x = ws.cell(i, c); x.font = Font(name=F); x.border = BORDER; x.fill = fill
            x.alignment = Alignment(horizontal='center' if c <= nk + 1 else 'right')
        for c in codes:
            sym = REGION_DEFAULTS[c]['symbol']
            ws.cell(i, col_of[(c, 'list')]).number_format = f'"{sym}"#,##0'
            if (c, 'taxed') in col_of: ws.cell(i, col_of[(c, 'taxed')]).number_format = f'"{sym}"#,##0.00'
            if (c, 'rmb') in col_of: ws.cell(i, col_of[(c, 'rmb')]).number_format = '¥#,##0'
            if (c, 'diff') in col_of: ws.cell(i, col_of[(c, 'diff')]).number_format = '¥#,##0;(¥#,##0);-'
            if (c, 'prem') in col_of: ws.cell(i, col_of[(c, 'prem')]).number_format = '0.0%'
    widths = [6] + [k.get('width', 12) for k in keys] + [15] * (len(heads) - nk - 1)
    for c, w in enumerate(widths, 1): ws.column_dimensions[L(c)].width = w
    ws.freeze_panes = ws.cell(2, nk + 2).coordinate
    ws.auto_filter.ref = f'A1:{L(len(heads))}{len(allkeys) + 1}'
    for c in codes:
        d = REGION_DEFAULTS[c]
        ws.cell(1, col_of[(c, 'list')]).comment = Comment(f'{d["label"]}官网定制页汇总价，{"含税" if d["tax_inclusive"] else "不含税"}。{d["note"]}', 'apple-official-price')
        if (c, 'taxed') in col_of: ws.cell(1, col_of[(c, 'taxed')]).comment = Comment(f'= 官网价 × (1 + 「{ASSUME}」D{region_row(c)} 税率)', 'apple-official-price')
        if (c, 'rmb') in col_of: ws.cell(1, col_of[(c, 'rmb')]).comment = Comment(f'= 含税价 × 「{ASSUME}」E{region_row(c)} 汇率（占位估算值，可改）', 'apple-official-price')

    # ---- add-ons ----
    addons = spec.get('addons', [])
    if addons:
        r = len(allkeys) + 4
        ws.cell(r, 1, '附加项差价（不进上表乘积；在任意配置上加选即加此差价）').font = Font(name=F, bold=True, color='1F3864')
        r += 1
        ah = ['附加项'] + [f'{REGION_DEFAULTS[c]["label"]}差价 ({REGION_DEFAULTS[c]["currency"]})' for c in codes] + ['备注']
        for c, h in enumerate(ah, 1):
            x = ws.cell(r, c, h); x.fill = HDR_FILL; x.font = HDR_FONT; x.border = BORDER; x.alignment = Alignment(horizontal='center')
        for a in addons:
            r += 1; ws.cell(r, 1, a['name'])
            for j, c in enumerate(codes, 2):
                v = a.get(c); x = ws.cell(r, j, v if v is not None else '未取到')
                x.number_format = f'"{REGION_DEFAULTS[c]["symbol"]}"#,##0;("{REGION_DEFAULTS[c]["symbol"]}"#,##0);-'
            ws.cell(r, len(codes) + 2, a.get('note', ''))
            for c in range(1, len(codes) + 3): ws.cell(r, c).font = Font(name=F); ws.cell(r, c).border = BORDER

    append_notes(ws_a, spec.get('model', spec['sheet']), spec)
    # keep assumptions sheet last
    wb._sheets = [s for s in wb._sheets if s.title != ASSUME] + [ws_a]
    wb.calculation.fullCalcOnLoad = True
    wb.save(spec['workbook'])
    print('saved', spec['workbook'], 'sheet', spec['sheet'], 'rows', len(allkeys), 'sheets', wb.sheetnames)

if __name__ == '__main__':
    main(sys.argv[1])
